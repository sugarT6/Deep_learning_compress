#!/usr/bin/env python3
"""Report quality-value distributions for metadata-registered FASTQ datasets."""

from __future__ import annotations

import argparse
import gzip
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Iterable

import numpy as np
from openpyxl import load_workbook


QUALITY_ASCII_OFFSET = 33
QUALITY_ALPHABET_SIZE = 95
QUALITY_CHUNK_BYTES = 8 * 1024 * 1024
DEFAULT_MAX_READS = 500_000
DEFAULT_SPECIES = "Homo sapiens"
DEFAULT_DATA_DIR = Path(__file__).resolve().parent / "data" / "2nd"
DEFAULT_DETAILS_FILENAME = "sequencing_platform_details.xlsx"


@dataclass(frozen=True)
class DatasetMetadata:
    accession: str
    platform: str
    species: str


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_dataset_metadata(path: Path) -> list[DatasetMetadata]:
    """Load per-accession species and platform labels from the details workbook."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["sequencing_platform_details"]
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_header = next(rows)
        except StopIteration as exc:
            raise ValueError(f"{path}: workbook sheet is empty") from exc

        header = {_cell_text(value): index for index, value in enumerate(raw_header)}
        required = {"Accession", "Platform", "Instrument_Model", "Species"}
        missing = sorted(required - header.keys())
        if missing:
            raise ValueError(f"{path}: missing columns: {', '.join(missing)}")

        records: list[DatasetMetadata] = []
        seen_accessions: set[str] = set()
        for row_number, row in enumerate(rows, start=2):
            accession = _cell_text(row[header["Accession"]])
            if not accession:
                continue
            if accession in seen_accessions:
                raise ValueError(
                    f"{path}: duplicate accession {accession!r} at row {row_number}"
                )

            platform_group = ""
            if "Platform_Group" in header:
                platform_group = _cell_text(row[header["Platform_Group"]])
            platform = (
                platform_group
                or _cell_text(row[header["Instrument_Model"]])
                or _cell_text(row[header["Platform"]])
            )
            species = _cell_text(row[header["Species"]])
            if not platform:
                raise ValueError(
                    f"{path}: accession {accession!r} has no platform label"
                )

            records.append(DatasetMetadata(accession, platform, species))
            seen_accessions.add(accession)
        return records
    finally:
        workbook.close()


def _query_matches_accession(query: str, accession: str) -> bool:
    name = Path(query).name.upper()
    accession_upper = accession.upper()
    return name == accession_upper or any(
        name.startswith(accession_upper + delimiter) for delimiter in ("_", ".")
    )


def select_datasets(
    records: Iterable[DatasetMetadata],
    datasets: Iterable[str] | None = None,
) -> list[DatasetMetadata]:
    """Select all human accessions by default, or any requested accessions."""

    records = list(records)
    datasets = list(datasets or [])
    if not datasets:
        selected = [
            record
            for record in records
            if record.species.casefold() == DEFAULT_SPECIES.casefold()
        ]
        if not selected:
            raise ValueError(f"metadata contains no {DEFAULT_SPECIES} datasets")
        return selected

    selected = []
    selected_accessions: set[str] = set()
    for dataset in datasets:
        matches = [
            record
            for record in records
            if _query_matches_accession(dataset, record.accession)
        ]
        if not matches:
            raise ValueError(
                f"dataset {dataset!r} is not present in the details workbook"
            )
        if len(matches) > 1:
            accessions = ", ".join(record.accession for record in matches)
            raise ValueError(f"dataset query {dataset!r} is ambiguous: {accessions}")

        match = matches[0]
        if match.accession not in selected_accessions:
            selected.append(match)
            selected_accessions.add(match.accession)
    return selected


def resolve_fastq_path(data_dir: Path, accession: str) -> Path:
    """Resolve exactly one FASTQ or gzipped FASTQ belonging to an accession."""

    candidates = []
    for path in sorted(data_dir.iterdir()):
        if not path.is_file() or not (
            path.name.endswith(".fastq.gz")
            or path.name.endswith(".fq.gz")
            or path.name.endswith(".fastq")
            or path.name.endswith(".fq")
        ):
            continue
        if _query_matches_accession(path.name, accession):
            candidates.append(path)

    if not candidates:
        raise FileNotFoundError(
            f"{data_dir}: no FASTQ file found for accession {accession}"
        )
    if len(candidates) > 1:
        names = ", ".join(path.name for path in candidates)
        raise ValueError(
            f"{data_dir}: multiple FASTQ files found for accession "
            f"{accession}: {names}"
        )
    return candidates[0]


def _open_fastq(path: Path) -> BinaryIO:
    if path.name.endswith(".gz"):
        return gzip.open(path, "rb")
    return path.open("rb")


def _without_line_ending(line: bytes) -> bytes:
    return line.rstrip(b"\r\n")


def count_fastq_quality_ids(
    path: Path,
    *,
    max_reads: int = DEFAULT_MAX_READS,
) -> np.ndarray:
    """Count Phred+33 ids from a strict four-line FASTQ without loading it all."""

    if max_reads <= 0:
        raise ValueError("max_reads must be positive")

    counts = np.zeros(QUALITY_ALPHABET_SIZE, dtype=np.int64)
    quality_chunk = bytearray()

    def flush_quality_chunk(last_record_number: int) -> None:
        if not quality_chunk:
            return
        raw_values = np.frombuffer(quality_chunk, dtype=np.uint8)
        minimum = int(raw_values.min())
        maximum = int(raw_values.max())
        maximum_allowed = QUALITY_ASCII_OFFSET + QUALITY_ALPHABET_SIZE - 1
        if minimum < QUALITY_ASCII_OFFSET or maximum > maximum_allowed:
            raise ValueError(
                f"{path}: quality data through record {last_record_number} contains "
                f"a byte outside Phred+33 ids Q0-Q{QUALITY_ALPHABET_SIZE - 1}"
            )
        raw_counts = np.bincount(
            raw_values, minlength=QUALITY_ASCII_OFFSET + QUALITY_ALPHABET_SIZE
        )
        counts[:] += raw_counts[
            QUALITY_ASCII_OFFSET : QUALITY_ASCII_OFFSET + QUALITY_ALPHABET_SIZE
        ]
        del raw_values
        quality_chunk.clear()

    with _open_fastq(path) as handle:
        record_number = 0
        while True:
            header = handle.readline()
            if not header:
                break
            sequence = handle.readline()
            plus = handle.readline()
            quality = handle.readline()
            record_number += 1
            if not sequence or not plus or not quality:
                raise ValueError(
                    f"{path}: truncated FASTQ record {record_number}"
                )

            header = _without_line_ending(header)
            sequence = _without_line_ending(sequence)
            plus = _without_line_ending(plus)
            quality = _without_line_ending(quality)
            if not header.startswith(b"@"):
                raise ValueError(
                    f"{path}: record {record_number} header does not start with @"
                )
            if not plus.startswith(b"+"):
                raise ValueError(
                    f"{path}: record {record_number} plus line does not start with +"
                )
            if len(sequence) != len(quality):
                raise ValueError(
                    f"{path}: record {record_number} sequence/quality lengths differ "
                    f"({len(sequence)} != {len(quality)})"
                )

            quality_chunk.extend(quality)
            if len(quality_chunk) >= QUALITY_CHUNK_BYTES:
                flush_quality_chunk(record_number)
            if record_number >= max_reads:
                break

        flush_quality_chunk(record_number)

    if not counts.sum():
        raise ValueError(f"{path}: FASTQ contains no quality values")
    return counts


def format_quality_distribution(
    counts: np.ndarray,
    *,
    entries_per_line: int = 6,
) -> str:
    """Format nonzero ``Q<id> percentage`` entries in ascending id order."""

    if entries_per_line <= 0:
        raise ValueError("entries_per_line must be positive")
    if counts.ndim != 1 or np.any(counts < 0):
        raise ValueError("counts must be a one-dimensional nonnegative array")
    total = int(counts.sum())
    if total == 0:
        raise ValueError("dataset contains no quality values")

    entries = [
        f"Q{quality_id} {100.0 * int(count) / total:.4f}%"
        for quality_id, count in enumerate(counts)
        if count
    ]
    return "\n".join(
        "\t".join(entries[start : start + entries_per_line])
        for start in range(0, len(entries), entries_per_line)
    )


def format_dataset_report(metadata: DatasetMetadata, counts: np.ndarray) -> str:
    """Format a platform/accession title followed by six entries per line."""

    title = f"{metadata.platform}\t{metadata.accession}"
    return f"{title}\n{format_quality_distribution(counts, entries_per_line=6)}"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Report FASTQ quality-value distributions for datasets listed in "
            "sequencing_platform_details.xlsx; the default selection is Homo sapiens."
        )
    )
    parser.add_argument(
        "datasets",
        nargs="*",
        help=(
            "accessions or FASTQ filenames to query, regardless of species; "
            "omit to query every Homo sapiens dataset"
        ),
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=DEFAULT_DATA_DIR,
        help=f"FASTQ directory; default: {DEFAULT_DATA_DIR}",
    )
    parser.add_argument(
        "--details-xlsx",
        type=Path,
        help=(
            "details workbook; default: "
            "<data-dir>/sequencing_platform_details.xlsx"
        ),
    )
    parser.add_argument(
        "--quiet",
        action="store_true",
        help="do not print per-dataset progress to stderr",
    )
    parser.add_argument(
        "--max-reads",
        type=int,
        default=DEFAULT_MAX_READS,
        help=(
            "maximum FASTQ records (quality lines) counted per dataset; "
            f"default: {DEFAULT_MAX_READS}"
        ),
    )
    return parser


def main() -> int:
    args = build_parser().parse_args()
    details_path = args.details_xlsx or args.data_dir / DEFAULT_DETAILS_FILENAME
    try:
        metadata = load_dataset_metadata(details_path)
        selected = select_datasets(metadata, args.datasets)
        reports = []
        for index, record in enumerate(selected, start=1):
            path = resolve_fastq_path(args.data_dir, record.accession)
            if not args.quiet:
                print(
                    f"[{index}/{len(selected)}] counting {record.accession}: "
                    f"{path.name}",
                    file=sys.stderr,
                    flush=True,
                )
            reports.append(
                format_dataset_report(
                    record,
                    count_fastq_quality_ids(path, max_reads=args.max_reads),
                )
            )
        print("\n\n".join(reports))
    except (FileNotFoundError, KeyError, OSError, ValueError) as exc:
        raise SystemExit(str(exc)) from exc
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
